from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
text = (ROOT / 'AI Usage Tracker Strategy Report.md').read_text()
definitions = dict(re.findall(r'^\[\^(\d+)\]: (.*)$', text, re.M))
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.72)
section.left_margin = section.right_margin = Inches(.78)
section.page_width = Inches(8.5)
section.page_height = Inches(11)
for style in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    doc.styles[style].font.name = 'Calibri'
    doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.widow_control = True
normal.paragraph_format.keep_together = True
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].paragraph_format.space_after = Pt(14)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
doc.core_properties.title = 'AI Usage Tracker Market and Product Strategy'
doc.core_properties.subject = 'Market research and product strategy'
doc.core_properties.author = ''

footroot = OxmlElement('w:footnotes')
for fid, typ in [(-1, 'separator'), (0, 'continuationSeparator')]:
    f = OxmlElement('w:footnote'); f.set(qn('w:id'), str(fid)); f.set(qn('w:type'), typ)
    p = OxmlElement('w:p'); r = OxmlElement('w:r'); r.append(OxmlElement('w:' + typ)); p.append(r); f.append(p); footroot.append(f)
footpart = Part(PackURI('/word/footnotes.xml'), 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml', b'', doc.part.package)
doc.part.relate_to(footpart, RT.FOOTNOTES)
fn_count = 0

def link(parent, part, label, url, size=None):
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '245579'); pr.append(color)
    if size:
        s = OxmlElement('w:sz'); s.set(qn('w:val'), str(size)); pr.append(s)
    r.append(pr); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); parent.append(h)

def footnote(p, source):
    global fn_count
    fn_count += 1
    r = p.add_run(); r.font.superscript = True
    ref = OxmlElement('w:footnoteReference'); ref.set(qn('w:id'), str(fn_count)); r._r.append(ref)
    f = OxmlElement('w:footnote'); f.set(qn('w:id'), str(fn_count))
    fp = OxmlElement('w:p'); pp = OxmlElement('w:pPr')
    sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '30'); pp.append(sp); fp.append(pp)
    fr = OxmlElement('w:r'); pr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '16'); pr.append(sz); fr.append(pr)
    fr.append(OxmlElement('w:footnoteRef')); fp.append(fr)
    spacer = OxmlElement('w:r'); st = OxmlElement('w:t'); st.set(qn('xml:space'), 'preserve'); st.text = ' '; spacer.append(st); fp.append(spacer)
    desc = definitions[source]
    m = re.search(r'\[([^\]]+)\]\(([^)]+)\)', desc)
    label = f' Source {source}: ' + (m.group(1) if m else desc)
    link(fp, footpart, label, m.group(2) if m else '', 16)
    rr = OxmlElement('w:r'); rp = OxmlElement('w:rPr'); ss = OxmlElement('w:sz'); ss.set(qn('w:val'), '16'); rp.append(ss); rr.append(rp)
    tt = OxmlElement('w:t'); tt.text = '. Accessed September 13, 2026.'; rr.append(tt); fp.append(rr)
    f.append(fp); footroot.append(f)

pattern = re.compile(r'(\[\^\d+\]|\*\*.*?\*\*|\[[^\]]+\]\([^)]+\))')
def inline(p, s, citations=True):
    s = s.replace('][^', '],\u00a0[^')
    for token in pattern.split(s):
        if not token: continue
        if re.fullmatch(r'\[\^\d+\]', token):
            if citations: footnote(p, re.search(r'\d+', token).group())
        elif token.startswith('**'):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith('['):
            m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
            link(p._p, doc.part, m.group(1), m.group(2))
        else: p.add_run(token)

lines = text.splitlines(); i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line.startswith('[^'):
        sid = re.match(r'\[\^(\d+)\]', line).group(1)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
        p.add_run(sid + '. ').bold = True
        inline(p, definitions[sid], False)
        for r in p.runs: r.font.size = Pt(9)
        i += 1; continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title'); i += 1; continue
    if line.startswith('## '):
        p = doc.add_paragraph(line[3:], 'Heading 1')
        if line.startswith('## 2 '): p.paragraph_format.page_break_before = True
        i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            if not re.match(r'^\|[\s:|\-]+\|$', lines[i]): rows.append([v.strip() for v in lines[i].strip('|').split('|')])
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0])); table.autofit = False
        widths = [1.35, 2.85, 2.74] if len(rows[0]) == 3 else [1.5, 5.44]
        if rows[0][0] == 'Evidence and status': widths = [2.2, 2.25, 2.49]
        for c,w in zip(table.columns,widths): c.width = Inches(w)
        for ri, vals in enumerate(rows):
            cells = table.add_row().cells
            for ci, val in enumerate(vals):
                cells[ci].width = Inches(widths[ci])
                p = cells[ci].paragraphs[0]; p.paragraph_format.space_after = Pt(5); p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.line_spacing = 1.0
                inline(p, val)
                for r in p.runs:
                    r.font.size = Pt(9)
                    if ri == 0: r.bold = True
                tcpr = cells[ci]._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'E7E9EC' if ri == 0 else ('F5F6F7' if ri % 2 else 'FFFFFF')); tcpr.append(shd)
                margins = OxmlElement('w:tcMar')
                for name in ['left', 'right']:
                    el = OxmlElement('w:' + name); el.set(qn('w:w'), '80'); el.set(qn('w:type'), 'dxa'); margins.append(el)
                tcpr.append(margins)
            trpr = table.rows[-1]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        continue
    p = doc.add_paragraph(); inline(p, line); i += 1

footpart._blob = etree.tostring(footroot, xml_declaration=True, encoding='UTF-8', standalone=True)
output = ROOT / 'AI Usage Tracker Strategy Report.docx'
doc.save(output)
print(f'Created {output}; {fn_count} footnotes; {len(text.split())} words including sources.')
